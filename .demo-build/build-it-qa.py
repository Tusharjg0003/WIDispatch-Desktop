from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/widispatch-it-20260929/WIDispatch_Client_IT_Questions_and_Answers.docx'

# Each answer is deliberately speakable. Follow-up notes are for the presenter.
sections = [
('Platform purpose and architecture', [
('What problem does WIDispatch solve',
 'WIDispatch supports coordinated water dispatch planning across production, transmission, storage and demand. It brings operational inputs and approved cost assumptions into a network model, calculates allocations and shortages, and lets planners review and publish decisions. The objective is an explainable operating plan that meets demand as far as available capacity allows and makes the consequences of constraints visible.',
 'Show the operating cycle, then one example linking approved demand to a resulting allocation. Avoid promising that every scenario can meet all demand.', 'D3 D5 D4'),
('Why are there five applications',
 'The four portals organise work by operational responsibility. Production records supply and operating information; Demand records city-gate requirements and receipts; Transmission records network availability and water movement; Economics manages cost information. Desktop is the planner’s workspace for the shared asset registry, network, scenarios, results and reviewed dispatch decisions. Input approval, simulation and publication are separate steps.',
 'Explain who owns each step. Published production requests, demand decisions, maintenance metadata and actual operating records have different meanings.', 'D3 D4 D6 P4 P8'),
('Is Desktop an installed Windows application',
 'In the implementation being demonstrated, Desktop is a browser-based planning application. Its interface uses React and Vite, with a separate Express and Node.js API. An installed Windows package is not established by the inspected source. The delivery method and supported browsers should be confirmed for the client environment.',
 'Use “planning workspace” when introducing Desktop. Do not promise an offline installer or local database on every user’s computer.', 'D1'),
('What is the architecture and how do the applications communicate',
 'Each portal uses Next.js with server-side API routes. Desktop has its own Express API. These server components use MongoDB collections and common asset identifiers to exchange operational and planning information. The current architecture relies on shared data contracts; the inspected implementation does not establish an event bus or a complete independent microservice boundary for each domain.',
 'Open the architecture diagram in the IT briefing. Distinguish browser traffic, server APIs and database access, and label the proposed hosting boundary.', 'D1 D2 D4')]),
('Hosting and deployment', [
('Can it run in our data centre or approved cloud',
 'The application components can be assessed for your approved hosting environment, including a client-managed environment. The final deployment design still needs agreement on hosting location, network access, identity, database operations and support ownership. The deployment diagram in this pack is a proposal, not evidence that a production environment has already been deployed and accepted.',
 'Ask for the approved operating system, cloud or data centre, network zones, database service and security requirements before committing to a topology.', 'D1 D2 R1'),
('What hardware and capacity do we need',
 'Production sizing needs a benchmark using your expected users, network size, planning horizon and data volume. The draft specification proposes a trial host with 4 virtual CPUs, 16 GB RAM and about 100 GB storage. Those figures are a starting point for a controlled trial, not a measured production recommendation or capacity guarantee.',
 'Ask for concurrent planners, daily records, retained history and maximum scenario size. Record the agreed workload before quoting performance.', 'R1'),
('Can users work offline or without internet access',
 'A complete offline mode has not been established in the inspected implementation. The browser applications need access to their services and database. The map views also reference external tile providers. A restricted network deployment would need approved access to those providers or a designed internal map alternative, plus tests of the required workflows without public internet access.',
 'Separate “no public internet” from “no connection to application servers.” Neither file import nor a browser cache proves offline operation.', 'D1 G1'),
('What must client IT configure before a trial',
 'We need agreed service URLs, a protected ingress, certificates, application and database configuration, secrets, test accounts, and backup and monitoring responsibilities. We also need to reconcile database naming across the applications. Some readiness items require code changes, particularly Desktop access control and publication integrity, so the trial plan should explicitly identify its permitted users and purpose.',
 'Use the Deployment and Readiness register tabs to assign owners. A protected network perimeter does not by itself implement application permissions.', 'D1 D2 R1')]),
('Identity roles and permissions', [
('Can we use corporate SSO and MFA',
 'The current portal login uses local accounts, password verification and signed tokens. Corporate single sign-on and multi-factor authentication integration were not identified in the inspected application paths. Federation with your identity provider is a development and acceptance item, including group mapping, logout and account revocation. Desktop API authentication also needs implementation.',
 'Ask which identity provider and protocol the client requires. Do not describe SSO as a setting that can simply be switched on today.', 'P2 D1 S1'),
('Which roles exist and who can approve',
 'The portals define GM, Regional Head, City or Governorate Head, Asset Manager, Read-only and Dashboard Approver roles. The inspected helper rules allow GM and Dashboard Approver approval actions, with operating scope and administrative rights varying by role. The matrix documents those rules and their exceptions. Separate Desktop planner and publisher permissions remain a proposed control, rather than an enforced API role boundary.',
 'Use the Roles tab for exact differences. A portal GM role does not grant operating-system or database administration rights.', 'P1 P3 D1'),
('Can one region see or approve another region’s records',
 'Scope helpers exist, but uniform isolation across every route is not yet proven. The source review found approval routes that check role and retrieve a record by ID without an explicit asset-scope check. It also found differences between portals and some asset-scope helper behaviour. These gaps need correction and cross-scope denial tests before we claim complete regional isolation or enforced separation of duties.',
 'Confirm whether self-approval is prohibited. A successful dashboard demonstration is not proof that direct API requests cannot cross a scope boundary.', 'P1 P3'),
('How are users provisioned and access revoked',
 'Portal user-management and assignment paths exist, but we need to validate the complete joiner, mover and leaver process for the client. The current token configuration uses a seven-day expiry. Immediate revocation and corporate directory-driven provisioning were not established by this review. Role changes, account disablement and outstanding sessions must be tested against the agreed access policy.',
 'Agree account ownership, approver provisioning and leaver response time. Do not claim that disabling an account immediately invalidates every issued token.', 'P1 P2 D2')]),
('Data ownership and input approval', [
('Which application owns the master data',
 'The proposed operating model places shared asset stewardship and network authoring in Desktop. Domain teams own their operational inputs, and Economics owns cost proposals and approvals. Stable asset identifiers connect the applications. Some portal asset routes also exist, so authoritative ownership and permitted edit paths need explicit agreement; the current code alone does not establish exclusive stewardship.',
 'Show the Data ownership tab. Agree who may create, rename, retire or delete an asset, including protection for assets referenced by saved plans.', 'D6 P8'),
('Is there one database and can environments be separated',
 'The applications are intended to share MongoDB-backed operational data, but the reviewed database selection is inconsistent. Some services use a fixed database name, others derive it from the connection URI, and some assignment routes use a different fixed name. Live connection settings were not inspected. We need to standardise this policy and prove environment separation with a cross-application test.',
 'Do not reveal connection strings. Demonstrate that a test asset and account appear only in the intended environment after configuration is reconciled.', 'D2'),
('Does simulation use only approved information',
 'Approval eligibility depends on the type of information. Approved demand provides the demand basis, and approved effective-dated financial entries can provide variable operating cost. Maintenance availability uses a broader set of eligible workflow states, while maintenance decision candidates require approval provenance. We should show those rules explicitly instead of claiming that every input is treated identically.',
 'Agree the client’s rule for submitted or revised maintenance. Input approval does not itself publish a dispatch plan.', 'D5 P8'),
('How are data quality history and retention handled',
 'There are application validations, approval records and saved dispatch plans, with lineage fields linking published decisions to a plan. These provide useful traceability. A complete immutable input snapshot, protected long-term audit retention and reference-safe asset deletion were not established in this review. Retention, data correction and evidence preservation need an agreed policy and corresponding controls.',
 'Ask for units, missing-data rules, retention periods and correction authority. Validate duplicates and imported data using client examples.', 'D3 D4 D6 P6 P7')]),
('Simulation method and confidence', [
('What exactly does the simulation optimise',
 'The solver uses daily sequential minimum-cost maximum-flow calculations with storage carried forward between days. It allocates available supply through the modelled network and considers variable operating cost. It should not be presented as pressure or head-loss hydraulic simulation, a proven global optimum across the full planning horizon, or total lifecycle-cost optimisation.',
 'Use a simple network with one constrained line to explain the model. State the modelling assumptions before interpreting results.', 'D5 T1'),
('How does it handle shortages outages and storage',
 'The calculations apply modelled capacities and availability, then show where available supply or network constraints prevent demand from being met. Storage carries between daily calculations according to the scenario configuration. Existing tests cover shortage, bottleneck, outage and storage behaviour. The plan remains dependent on accurate asset data, topology, assumptions and the eligibility rules for operational inputs.',
 'Show one baseline and one constrained scenario. Avoid treating a mathematically feasible flow as proof of field pressure, water quality or every operating restriction.', 'D3 D5 T1'),
('Which costs are included and what happens if cost data is missing',
 'Dispatch uses the latest applicable approved plant variable operating-cost entry when available. The reviewed fallback is the plant’s variable O&M specification, followed by a coded default of 2.10 SAR per cubic metre. The client should approve the fallback policy or require incomplete-cost scenarios to be blocked. Fixed and capital cost information in Economics does not mean the solver optimises lifecycle cost.',
 'Confirm currency, units, effective dates and fallback visibility. Do not describe the coded default as a client-approved assumption.', 'D5 P8'),
('How do we know the results are correct',
 'The existing backend suite was rerun on 29 September 2026, with 87 tests passing and none failing or skipped. It covers solver and supporting logic, including capacities, storage, cost rules and decision validation. That is useful functional evidence. Client acceptance still needs agreed datasets, expected results, operational review and complete testing across all five applications.',
 'Show the Evidence tab and retain the test log. Unit tests do not establish a production security assessment, load benchmark or live database acceptance.', 'T1 R1')]),
('Reviewing decisions and publishing', [
('Can planners edit the recommended decisions',
 'Supported demand decisions can be edited within quantity bounds, and maintenance decisions can be reviewed. A material limitation is that these edits do not recalculate network flow or plant allocations. An edited decision should therefore not be described as a newly optimised feasible plan. We need an agreed revalidation process and, where required, implementation of recalculation or consistency checks.',
 'For an operating assumption change, create or rerun an appropriate scenario and review the resulting plan. Do not imply that editing a decision automatically triggers that run.', 'D3 D4 T1'),
('Are comments compulsory and can a published plan be changed',
 'A reduced demand decision requires a comment, and a rejected maintenance decision requires a comment. Demand approval quantities must be finite and within zero to the required quantity. The inspected routes reject edits to a published plan. Those controls support review accountability, but authenticated Desktop actor attribution still needs access-control work.',
 'Demonstrate one required-comment validation and one published-plan edit guard. Do not say that every decision requires a comment.', 'D3 D1 T1'),
('What if inputs change or two planners edit at once',
 'The reviewed path does not establish a version check tying publication to unchanged input records, or a revision check that prevents two draft editors overwriting one another. The draft-status guard is not a complete concurrency control. Versioned inputs, conflict detection and a clear revalidation policy are development items for a controlled production workflow.',
 'For a trial, use an agreed single-editor procedure and record the input cut-off. A procedure reduces exposure but does not replace the missing technical controls.', 'D3 D4'),
('Does pressing Run immediately send instructions to the portals',
 'Run creates a saved draft plan and results for review. Publication is a separate, explicit action after the planner reviews the decisions and required comments. Publication writes defined planning and decision fields back to operational collections. It does not itself send commands to pumps or valves, or convert planned volumes into measured actuals.',
 'Pause at the draft state during the demonstration so the client can see the distinction between calculation, review and publication.', 'D3 D4')]),
('Publication mapping and failure handling', [
('What data changes when a plan is published',
 'Production allocations are written to the production requested-volume field, creating a draft record if necessary. Demand decisions update Desktop-approved volume and decision metadata on existing demand records. Maintenance decisions update decision metadata without rescheduling the maintenance dates. Plan identifiers provide lineage. Actual receipt and delivery records and approved financial entries are not written by this publication path.',
 'Use the Publication tab for exact field names and matching rules. Existing actual values and portal submission or approval states are preserved.', 'D4 P4'),
('Will every portal display every published decision',
 'We should demonstrate each confirmed mapping separately. The Production daily table uses the requested allocation field. The inspected Demand daily table reads required demand and actual receipts rather than the Desktop-approved volume field, so we should not promise that view displays the dispatch allocation. Transmission system visibility follows saved registry and network data. Other decision surfaces need specific validation or UI changes.',
 'Keep the Desktop decision view available when demonstrating demand outcomes. Economics cost approval and transmission actual reporting remain separate workflows.', 'D4 D6 P4 P8'),
('What happens if publication fails halfway through',
 'The current implementation performs production, demand and maintenance writes in parallel and then marks the plan published. The inspected path does not wrap those writes in one database transaction, so a partial update is possible if a write fails. Transactional publication or a designed recoverable publication process, with fault and retry tests, is a development requirement.',
 'Do not say it is all-or-nothing today. Also reconcile missing target records; a successful operation alone does not prove every intended record was matched.', 'D3 D4'),
('Can we publish twice or undo a publication',
 'A later publication request for an already-published plan is rejected, and published-plan edits are blocked. That does not establish protection against simultaneous publication requests or safe retries after partial failure. A validated undo or superseding-plan workflow was not established in the inspected path. We need explicit versioning, recovery and correction rules before promising those behaviours.',
 'If asked about rollback, distinguish correcting published business data from restoring a database backup or rolling back application code.', 'D3 D4')]),
('Interfaces and external dependencies', [
('Which integrations are available today',
 'The source contains internal MongoDB-backed data sharing, supported spreadsheet exchanges, Desktop network file imports and exports, map services and email adapter paths. These have different configuration and testing needs. ERP, maintenance-management and other enterprise connectors were not identified in the inspected application paths. Each requested connector needs an agreed interface and implementation scope.',
 'Ask for the system, owner, data direction, identifiers, frequency, units and error-handling requirements. A file import is not a live synchronisation service.', 'D2 P5 P6 G1 S1'),
('Can it connect to SCADA or control equipment',
 'No SCADA or historian connector, or equipment-control connector, was identified in the inspected application paths. Telemetry ingestion would need a defined interface, aggregation rules and data-quality handling. Direct equipment control would be a separate scope with the client’s operational technology team. Publishing a dispatch plan currently updates application records rather than issuing actuator commands.',
 'Establish whether the request is read-only telemetry, operator instructions or automatic control; these require different designs and acceptance evidence.', 'D4 S1'),
('Can we use our GIS and internal maps',
 'Desktop includes network file exchange such as JSON and KMZ or KML import, with supported export paths. That can support a controlled exchange of network geometry and identifiers. It does not establish live synchronisation with an enterprise GIS. Existing map views reference external OpenStreetMap and ArcGIS imagery, so internal map services need compatibility work and validation.',
 'Ask for coordinate systems, asset IDs, geometry conventions, map-service endpoints and approved external-provider usage.', 'G1'),
('Are email notifications and central monitoring ready',
 'Email provider paths exist for SMTP, SendGrid and SES, with console behaviour as the default. Actual delivery was not tested in this review and requires approved configuration and any needed provider dependencies. Portal audit and service logging provide starting points for monitoring, but a complete SIEM integration and full Desktop actor coverage have not been established.',
 'Ask for the approved relay, recipients, event schema, log destination and alert owner. Protect or disable diagnostic email routes before production exposure.', 'P5 P7 D1 S1')]),
('Security evidence and production readiness', [
('Is the platform production ready and secure',
 'The operating workflow and functional evidence can be demonstrated, but production readiness has not been established by this review. The current source has material gaps, including Desktop API authentication, approval-scope enforcement, signing-secret fallback handling and publication integrity. The readiness register separates existing capabilities, configuration work and development work. Production approval should follow closure and testing of the agreed requirements.',
 'Answer directly and use the register. Do not treat an internal network, a successful login or 87 passing tests as proof of production security.', 'D1 P2 P3 D4 T1'),
('Is data encrypted and where will it be hosted',
 'Hosting region, transport encryption, database encryption and key custody need to be confirmed in the chosen deployment. The review did not inspect deployed certificates, reverse-proxy settings or database encryption configuration. We can define the required controls with your security team and provide deployment evidence once implemented and tested.',
 'Ask for residency, key ownership, backup location and approved data classifications. Do not assert compliance or a hosting jurisdiction without evidence.', 'D1 D2 R1'),
('Can we audit who changed and approved information',
 'Portal approval paths record actor and time information, and publication includes plan identifiers that support tracing decisions to a plan. That is a useful foundation, but it does not establish immutable audit retention or complete authenticated attribution for Desktop actions. Those controls need to be completed and tested against the client’s audit requirements.',
 'Demonstrate an approval record and a plan reference. Ask which events, before-and-after values, retention period and export format the client requires.', 'P3 P7 D4 D1'),
('Have you completed penetration testing or certification',
 'No penetration-test report, complete dependency-security assessment or certification evidence was reviewed for this pack. Source inspection and functional tests are different types of evidence. We should agree the required assessment scope, provide the relevant environment and documentation, resolve findings and retest before making a security or compliance claim.',
 'If asked for a report, record the exact report and standard required. Do not invent a certificate, test date, severity rating or remediation deadline.', 'S1 R1')]),
('Performance recovery and service ownership', [
('How many users and how large a network can it support',
 'A measured production capacity has not yet been established. The draft trial assumptions include five internal test accounts, one active run, 50 nodes, 100 edges and a 14-day horizon. These are proposed test conditions, not proven limits. We should benchmark your workload on the chosen host and report response times, run duration, resource use and behaviour under concurrent activity.',
 'The 87-test runner duration is not a simulation-speed measurement. Avoid quoting draft response targets as a service guarantee.', 'R1 T1'),
('What are the backup recovery and availability guarantees',
 'Backup configuration, restoration and failover have not been verified in this review. The draft trial proposes a recovery point of up to 24 hours, recovery time of up to 8 hours and 14-day retention, but these are not agreed service levels. Production targets, backup ownership and availability design need agreement, followed by a timed restore and failure exercise.',
 'Explain recovery point as acceptable data loss and recovery time as time to restore service. A single VM or single-node replica set does not prove high availability.', 'R1'),
('How will releases upgrades and rollback be controlled',
 'Build scripts exist, but a validated production release and rollback pipeline was not established in this review. The release process should include explicit checks, a recorded application version, compatible data changes, backups and a tested rollback procedure. The source review also found a build setting that can ignore TypeScript errors, so a build alone is not a sufficient quality gate.',
 'Ask for change windows, approval workflow and patching ownership. Treat application rollback and database or published-decision recovery as separate procedures.', 'S1 R1'),
('Who supports the system and what are the handover terms',
 'The support model and commercial terms need explicit agreement. We should define who operates the applications and database, handles incidents, approves changes and maintains integrations, together with support hours and escalation. Data ownership, retention, export scope and exit arrangements should also be documented. I would confirm those terms in writing rather than improvise a contractual commitment during the demonstration.',
 'Supported file exports are useful but do not prove a complete archive or migration facility. Bring the approved support and licensing proposal if one exists.', 'G1 R1')])
]

doc = Document()
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin, s.bottom_margin = Inches(.68), Inches(.65)
s.left_margin, s.right_margin = Inches(.78), Inches(.78)
s.header_distance, s.footer_distance = Inches(.28), Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.underline = False
    for color in st.element.xpath('.//w:color'):
        for attr in ['themeColor', 'themeTint', 'themeShade']:
            color.attrib.pop(qn('w:' + attr), None)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.07
normal.paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(28)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(13)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 2'].paragraph_format.keep_with_next = True

header = s.header.paragraphs[0]
header.text = 'WIDispatch    Client IT demonstration preparation'
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rr = footer.add_run('Presenter guide  |  29 September 2026  |  ')
rr.font.size = Pt(8)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)

def p(text, boldlead=None, size=None, after=None):
    para = doc.add_paragraph()
    if boldlead:
        para.add_run(boldlead + ': ').bold = True
    r = para.add_run(text)
    if size: r.font.size = Pt(size)
    if after is not None: para.paragraph_format.space_after = Pt(after)
    return para

def newpage(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

doc.add_paragraph('WIDispatch client IT questions and answers', 'Title')
doc.add_paragraph('Presenter preparation guide', 'Subtitle')
p('Use this guide to prepare for the technical discussion accompanying the Desktop and four-portal demonstration. The suggested answers are written to say aloud. The presenter notes identify what to show, what to confirm and where the current implementation needs work.')
p('The operating workflow has source and functional-test evidence. Production security, deployment capacity and recovery readiness still require specific implementation and acceptance evidence. Keep that distinction clear throughout the meeting.', boldlead='Main message')
doc.add_heading('Prepare these questions first', 2)
p('If you have limited time, rehearse Q3–Q4 on architecture, Q9–Q11 on identity and permissions, Q17 on the solver, Q21 on manual edits, Q25–Q27 on publication, Q33 on readiness and Q38 on recovery. These are the twelve priority answers.')
doc.add_heading('A short opening you can use', 2)
p('“Today I’ll show how WIDispatch takes operational information through planning, review and publication, and how the applications share responsibility for that cycle. I’ll also distinguish the capabilities in the current implementation from the configuration, integration and assurance work needed for your production environment.”')
doc.add_heading('How to use the guide', 2)
p('Give the suggested answer first, then pause. Use the presenter note only when the client asks for detail or evidence. Source codes point to the reference page. They are evidence pointers for your preparation, not independent certification.')
p('When an answer depends on the client’s requirements, say: “Let’s record the requirement, the evidence you need and the owner who will confirm it.” Avoid committing to dates, service levels or certifications that have not been agreed.')
doc.add_heading('Where to find each topic', 2)
for label, nums, page in [
    ('Purpose and architecture', '1–4', 2), ('Hosting and deployment', '5–8', 3),
    ('Identity and permissions', '9–12', 4), ('Data and approval', '13–16', 5),
    ('Simulation', '17–20', 6), ('Decision review', '21–24', 7),
    ('Publication', '25–28', 8), ('Interfaces', '29–32', 9),
    ('Security', '33–36', 10), ('Performance and service', '37–40', 11)]:
    p(f'{label}   ·   Q{nums}   ·   page {page}', size=10, after=2)

n=0
for title, questions in sections:
    newpage(title)
    for question, answer, note, refs in questions:
        n+=1
        doc.add_heading(f'Q{n}  {question}', 2)
        p(answer, boldlead='Suggested answer')
        para=p(note, boldlead='Presenter note', size=10)
        for run in para.runs: run.font.size=Pt(10)
        para=p(f'Evidence  {refs}', size=8, after=1)
        para.paragraph_format.keep_with_next=False

newpage('Preparing for the meeting')
doc.add_heading('Rehearse within one day', 2)
for text in [
 'First 45 minutes: read the twelve priority answers on page 1 aloud. Keep each initial answer under a minute; use the notes for follow-up detail.',
 'Next 60 minutes: rehearse the operating cycle and align each screen with its owner, input state and expected output. Keep a known baseline scenario available.',
 'Next 45 minutes: rehearse the architecture diagram, role matrix and exact publication mapping from the IT reference pack.',
 'Next 30 minutes: review the readiness register and identify who can confirm open security, hosting, integration and recovery items.',
 'Final 30 minutes: check the files and demonstration environment, prepare a screenshot fallback, and practise recording questions without promising an unagreed delivery date.'
]: p(text)
doc.add_heading('Questions to ask the client IT team', 2)
for text in [
 'What hosting location, network zones, data residency and database services are approved?',
 'Which identity provider, MFA policy, role hierarchy and approval separation must apply?',
 'What workloads, concurrent users and operating hours define acceptance?',
 'Which external systems need integration, and who owns their interfaces and asset identifiers?',
 'What security assessment, audit retention, backup, recovery and availability evidence is required?',
 'Who will operate and support the service, and who signs off the pilot and production acceptance?'
]: p(text)
doc.add_heading('If a demonstration step fails', 2)
p('“This step has not completed successfully in this environment. I’ll show the intended workflow using the prepared material and record the issue for verification. I won’t treat this as a completed transaction until we reconcile the saved plan and affected records.”')
p('For publication errors, avoid repeated clicks. Record the plan ID, time and visible error, then have engineering reconcile production, demand and maintenance updates. The current path can partially write data before failing.')
doc.add_heading('Keep a clear follow up record', 2)
p('Capture the client’s question, the exact requirement, evidence requested, responsible owner and an agreed response date. Separate an answer requiring deployment confirmation from a feature requiring development. Use the workbook register to track closure.')

newpage('Evidence and reference notes')
p('Basis: source review of the five supplied repositories and an existing backend test run on 29 September 2026. The review was targeted; it was not a penetration test, live configuration inspection or complete client acceptance test. The SRD is an internal draft dated 10 September 2026; its trial targets are proposals.')
p('Companion files: WIDispatch_IT_Briefing.pptx and WIDispatch_IT_Reference_and_Readiness.xlsx. The workbook holds the detailed role exceptions, field mappings, readiness actions and source references. Keep both available during preparation.')
refs = [
 ('D1', 'Desktop architecture and access boundary', 'backend/src/server.js; frontend/src/api/client.js; frontend and backend package manifests.'),
 ('D2', 'Database selection and assignment', 'Desktop backend/src/db.js; all portal lib/mongodb.ts; Production, Transmission and Economics app/api/users/assign/route.ts.'),
 ('D3', 'Run review and publication lifecycle', 'Desktop backend/src/simulationConfigs.js, including decision validation and plan status guards.'),
 ('D4', 'Publication fields and write order', 'Desktop backend/src/simulation/writeback.js.'),
 ('D5', 'Simulation and cost rules', 'Desktop backend/src/simulation/dispatch.js, capacity.js, cost.js, graph.js and mcmf.js.'),
 ('D6', 'Registry and networks', 'Desktop backend/src/assetRegistry.js, networks.js and transmissionRegistry.js.'),
 ('P1 P2 P3', 'Roles login and approval boundaries', 'All four portals: lib/rbac.ts, lib/types.ts, lib/auth.ts, app/api/auth/login/route.ts and app/api/submissions/[id]/status/route.ts.'),
 ('P4', 'Visible production and demand fields', 'Production and Demand components/production/production-input-table.tsx.'),
 ('P5 P6', 'Email and spreadsheet interfaces', 'Portal lib/email.ts and app/api/test-email/route.ts; supported lib/imports.ts and Excel import dialogs.'),
 ('P7 P8', 'Audit and financial inputs', 'Portal audit helpers and audit routes; Economics financial-entries API and lib/financial-costs.ts.'),
 ('G1', 'Network exchange and external maps', 'Desktop NetworkBuilderPage.jsx, lib/networkKmz.js, NetworkCanvasMapView.jsx and AssetDetailPage.jsx.'),
 ('T1', 'Executed functional evidence', 'Desktop backend test run on 29 September 2026: 87 passed, 0 failed, 0 skipped. Execution log retained with preparation evidence.'),
 ('R1 S1', 'Requirements and review scope', 'WIDispatch_SRD.docx internal draft v1.0; targeted source and dependency-manifest inventory across the five repositories.')
]
for code, title, body in refs:
    para=p(f'{title}. {body}', boldlead=code, size=9.5, after=7)
    for run in para.runs: run.font.size=Pt(9.5)
p('Reconfirm material answers after code or environment changes. “Not identified” means not found in the inspected paths; it is not an exhaustive claim about every possible external deployment component.', size=9)

doc.core_properties.title='WIDispatch client IT questions and answers'
doc.core_properties.subject='Presenter preparation for the Desktop and four portal demonstration'
doc.core_properties.author='WIDispatch'
doc.core_properties.keywords='WIDispatch, client IT, demonstration, questions and answers'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(f'Saved {OUT}; {n} questions; {len(doc.paragraphs)} paragraphs')
